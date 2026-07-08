#!/usr/bin/env python3
"""Build final MPU6050 experiment plots, DOCX, and PDF from collected CSV files."""

from __future__ import annotations

import csv
import math
import statistics
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image as PdfImage
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output" / "final_report"
PLOT_DIR = OUT_DIR / "plots"
DIAGRAM_OUT = OUT_DIR / "mpu6050_connection_diagram.png"
DOCX_OUT = ROOT / "output" / "docx" / "学号_姓名_MPU6050实验报告.docx"
PDF_OUT = ROOT / "output" / "pdf" / "学号_姓名.pdf"
PITCH_SIGN = -1.0

CSV_COLUMNS = [
    "ms",
    "state",
    "ax_raw",
    "ay_raw",
    "az_raw",
    "gx_raw",
    "gy_raw",
    "gz_raw",
    "ax_g",
    "ay_g",
    "az_g",
    "ax_ms2",
    "ay_ms2",
    "az_ms2",
    "gx_dps",
    "gy_dps",
    "gz_dps",
    "pitch_acc",
    "pitch_comp",
    "pitch_dmp",
    "roll_dmp",
    "yaw_dmp",
]

DATASETS = {
    "static": {"file": ROOT / "static.csv", "label": "静止"},
    "tilt": {"file": ROOT / "tilt.csv", "label": "前倾/后倾"},
    "shake": {"file": ROOT / "shake.csv", "label": "快速晃动"},
}

PLOT_TITLES = {
    "static": "Static - DMP vs Complementary Pitch",
    "tilt": "Tilt - DMP vs Complementary Pitch",
    "shake": "Shake - DMP vs Complementary Pitch",
}


def read_rows(path: Path, state: str) -> list[dict[str, float | str]]:
    rows: list[dict[str, float | str]] = []
    with path.open("r", encoding="utf-8-sig", errors="ignore", newline="") as handle:
        for line in handle:
            if not line.strip() or line.startswith("#") or line.startswith("ms,state,"):
                continue
            parts = next(csv.reader([line]))
            if len(parts) != len(CSV_COLUMNS):
                continue
            row: dict[str, float | str] = {"state": state}
            try:
                for index, column in enumerate(CSV_COLUMNS):
                    if column == "state":
                        continue
                    row[column] = float(parts[index])
            except ValueError:
                continue
            for column in ("pitch_acc", "pitch_comp", "pitch_dmp"):
                row[column] = float(row[column]) * PITCH_SIGN
            rows.append(row)
    return rows


def mean(values: list[float]) -> float:
    return statistics.fmean(values) if values else 0.0


def stdev(values: list[float]) -> float:
    return statistics.stdev(values) if len(values) > 1 else 0.0


def stats(rows: list[dict[str, float | str]], column: str) -> dict[str, float]:
    values = [float(row[column]) for row in rows]
    return {
        "min": min(values),
        "max": max(values),
        "mean": mean(values),
        "std": stdev(values),
    }


def median_period_ms(rows: list[dict[str, float | str]]) -> float:
    times = [float(row["ms"]) for row in rows]
    deltas = [b - a for a, b in zip(times, times[1:]) if b > a]
    return statistics.median(deltas) if deltas else 0.0


def load_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for font_path in [
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibri.ttf"),
    ]:
        if font_path.exists():
            return ImageFont.truetype(str(font_path), size)
    return ImageFont.load_default()


def draw_plot(rows: list[dict[str, float | str]], title: str, out_path: Path) -> None:
    width, height = 1100, 620
    margin_left, margin_right, margin_top, margin_bottom = 90, 35, 55, 85
    plot_w = width - margin_left - margin_right
    plot_h = height - margin_top - margin_bottom
    font = load_font(22)
    small = load_font(18)

    times = [(float(row["ms"]) - float(rows[0]["ms"])) / 1000.0 for row in rows]
    dmp = [float(row["pitch_dmp"]) for row in rows]
    comp = [float(row["pitch_comp"]) for row in rows]
    x_min, x_max = min(times), max(times)
    y_min = min(min(dmp), min(comp))
    y_max = max(max(dmp), max(comp))
    pad = max(2.0, (y_max - y_min) * 0.12)
    y_min -= pad
    y_max += pad

    def x_map(x: float) -> int:
        if x_max == x_min:
            return margin_left
        return int(margin_left + (x - x_min) / (x_max - x_min) * plot_w)

    def y_map(y: float) -> int:
        if y_max == y_min:
            return margin_top + plot_h // 2
        return int(margin_top + (y_max - y) / (y_max - y_min) * plot_h)

    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle([margin_left, margin_top, margin_left + plot_w, margin_top + plot_h], outline="black", width=2)
    draw.text((width // 2, 18), title, fill="black", font=font, anchor="ma")

    for i in range(6):
        x = margin_left + int(plot_w * i / 5)
        draw.line([(x, margin_top), (x, margin_top + plot_h)], fill=(210, 210, 210), width=1)
        value = x_min + (x_max - x_min) * i / 5
        draw.text((x, margin_top + plot_h + 10), f"{value:.1f}", fill="black", font=small, anchor="ma")
    for i in range(6):
        y = margin_top + int(plot_h * i / 5)
        draw.line([(margin_left, y), (margin_left + plot_w, y)], fill=(210, 210, 210), width=1)
        value = y_max - (y_max - y_min) * i / 5
        draw.text((margin_left - 10, y), f"{value:.0f}", fill="black", font=small, anchor="rm")

    dmp_points = [(x_map(x), y_map(y)) for x, y in zip(times, dmp)]
    comp_points = [(x_map(x), y_map(y)) for x, y in zip(times, comp)]
    if len(dmp_points) > 1:
        draw.line(dmp_points, fill="black", width=3)
    if len(comp_points) > 1:
        for start, end in zip(comp_points, comp_points[1:]):
            if ((start[0] // 14) % 2) == 0:
                draw.line([start, end], fill="black", width=2)

    draw.text((margin_left + plot_w // 2, height - 35), "Time (s)", fill="black", font=small, anchor="ma")
    draw.text((45, margin_top + plot_h // 2), "Pitch (deg)", fill="black", font=small, anchor="mm")
    legend_x, legend_y = width - 360, 75
    draw.line([(legend_x, legend_y), (legend_x + 70, legend_y)], fill="black", width=3)
    draw.text((legend_x + 85, legend_y), "DMP Pitch", fill="black", font=small, anchor="lm")
    for x in range(legend_x, legend_x + 70, 14):
        draw.line([(x, legend_y + 35), (x + 8, legend_y + 35)], fill="black", width=2)
    draw.text((legend_x + 85, legend_y + 35), "Complementary Pitch", fill="black", font=small, anchor="lm")
    image.save(out_path)


def draw_connection_diagram(out_path: Path) -> None:
    width, height = 1200, 560
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = load_font(28)
    label_font = load_font(22)
    small_font = load_font(18)

    stm = (90, 115, 430, 450)
    mpu = (770, 115, 1110, 450)
    draw.rounded_rectangle(stm, radius=18, outline="black", width=4)
    draw.rounded_rectangle(mpu, radius=18, outline="black", width=4)
    draw.text(((stm[0] + stm[2]) // 2, 150), "STM32F103", fill="black", font=title_font, anchor="ma")
    draw.text(((stm[0] + stm[2]) // 2, 188), "ByCar Main Board", fill="black", font=small_font, anchor="ma")
    draw.text(((stm[0] + stm[2]) // 2, 218), "USART1", fill="black", font=small_font, anchor="ma")
    draw.text(((stm[0] + stm[2]) // 2, 242), "PA9 TX / PA10 RX", fill="black", font=small_font, anchor="ma")
    draw.text(((mpu[0] + mpu[2]) // 2, 150), "MPU6050", fill="black", font=title_font, anchor="ma")
    draw.text(((mpu[0] + mpu[2]) // 2, 188), "6-axis IMU Module", fill="black", font=small_font, anchor="ma")

    connections = [
        ("PB14", "SCL", 270, "I2C Clock"),
        ("PB15", "SDA", 315, "I2C Data"),
        ("3V3", "VCC", 360, "Power"),
        ("GND", "GND", 405, "Ground"),
    ]
    optional = ("PB9", "INT", 445, "Optional Interrupt")

    for left, right, y, note in connections:
        draw.text((stm[2] - 35, y), left, fill="black", font=label_font, anchor="rm")
        draw.text((mpu[0] + 35, y), right, fill="black", font=label_font, anchor="lm")
        draw.line([(stm[2], y), (mpu[0], y)], fill="black", width=3)
        draw.polygon([(mpu[0] - 12, y - 6), (mpu[0], y), (mpu[0] - 12, y + 6)], outline="black", fill="white")
        draw.text((600, y - 12), note, fill="black", font=small_font, anchor="ma")

    left, right, y, note = optional
    draw.text((stm[2] - 35, y), left, fill="black", font=label_font, anchor="rm")
    draw.text((mpu[0] + 35, y), right, fill="black", font=label_font, anchor="lm")
    x = stm[2]
    while x < mpu[0]:
        draw.line([(x, y), (min(x + 18, mpu[0]), y)], fill="black", width=2)
        x += 34
    draw.text((600, y - 12), note, fill="black", font=small_font, anchor="ma")

    draw.text((width // 2, 40), "MPU6050 to STM32 Wiring Diagram", fill="black", font=title_font, anchor="ma")
    draw.text((width // 2, 515), "PB14/PB15: software I2C   3V3/GND: power reference   PB9 INT: optional", fill="black", font=small_font, anchor="ma")
    image.save(out_path)


def set_run(run, size: float = 10.5, bold: bool = False, font: str = "SimSun") -> None:
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = bold


def add_para(doc: Document, text: str = "", size: float = 10.5, bold: bool = False, align=None):
    para = doc.add_paragraph()
    if align is not None:
        para.alignment = align
    para.paragraph_format.space_after = Pt(5)
    run = para.add_run(text)
    set_run(run, size=size, bold=bold)
    return para


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    size = 16 if level == 1 else 13
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    para.paragraph_format.space_after = Pt(5)
    run = para.add_run(text)
    set_run(run, size=size, bold=True, font="SimHei")


def set_cell(cell, text: str, size: float = 8.5, bold: bool = False, align=WD_ALIGN_PARAGRAPH.CENTER) -> None:
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = align
    para.paragraph_format.space_after = Pt(0)
    run = para.add_run(text)
    set_run(run, size=size, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")


def add_table(doc: Document, rows: list[list[str]], widths: list[float], font_size: float = 8.5) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_i, row in enumerate(rows):
        for col_i, text in enumerate(row):
            cell = table.cell(row_i, col_i)
            cell.width = Inches(widths[col_i])
            set_cell(cell, text, size=font_size, bold=(row_i == 0))
    table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_docx(all_rows: dict[str, list[dict[str, float | str]]], plot_paths: dict[str, Path]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.75)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("MPU6050 传感器通信与姿态解算实验报告")
    set_run(run, size=18, bold=True, font="SimHei")
    add_para(doc, "文件名占位：学号_姓名.pdf / 学号_姓名.docx", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    add_heading(doc, "一、实验连接与采集设置")
    doc.add_picture(str(DIAGRAM_OUT), width=Inches(6.4))
    add_para(doc, "图：MPU6050 与 STM32 的 I2C、供电和串口连接。", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_table(
        doc,
        [
            ["STM32 引脚", "MPU6050 引脚", "说明"],
            ["PB14", "SCL", "软件 I2C 时钟线"],
            ["PB15", "SDA", "软件 I2C 数据线"],
            ["3V3", "VCC", "传感器供电"],
            ["GND", "GND", "公共地"],
            ["PA9", "USART1_TX", "串口输出 115200 8N1"],
            ["PB9", "INT", "可选中断输入"],
        ],
        [1.3, 1.5, 3.8],
    )
    add_para(doc, "采样周期约 20 ms，对应更新频率约 50 Hz。根据实测安装方向，车体前后倾对应 DMP 的 Roll 轴，报告中的 Pitch 已按统一符号校正处理。")

    add_heading(doc, "二、三种状态采集概况")
    overview = [["状态", "文件", "有效样本", "时长(s)", "周期(ms)", "频率(Hz)", "Pitch DMP 范围(deg)"]]
    for key, meta in DATASETS.items():
        rows = all_rows[key]
        period = median_period_ms(rows)
        duration = (float(rows[-1]["ms"]) - float(rows[0]["ms"])) / 1000.0
        pitch = stats(rows, "pitch_dmp")
        overview.append([
            meta["label"],
            meta["file"].name,
            str(len(rows)),
            f"{duration:.2f}",
            f"{period:.1f}",
            f"{1000.0 / period:.1f}" if period else "-",
            f"{pitch['min']:.2f} ~ {pitch['max']:.2f}",
        ])
    add_table(doc, overview, [0.85, 1.05, 0.8, 0.75, 0.75, 0.75, 1.55], font_size=8)

    doc.add_page_break()
    add_heading(doc, "三、静止原始数据记录")
    static_rows = all_rows["static"][:10]
    raw_table = [["序号", "ms", "ax_raw", "ay_raw", "az_raw", "gx_raw", "gy_raw", "gz_raw", "pitch_dmp"]]
    for i, row in enumerate(static_rows, 1):
        raw_table.append([
            str(i),
            f"{float(row['ms']):.0f}",
            f"{float(row['ax_raw']):.0f}",
            f"{float(row['ay_raw']):.0f}",
            f"{float(row['az_raw']):.0f}",
            f"{float(row['gx_raw']):.0f}",
            f"{float(row['gy_raw']):.0f}",
            f"{float(row['gz_raw']):.0f}",
            f"{float(row['pitch_dmp']):.3f}",
        ])
    add_table(doc, raw_table, [0.45, 0.75, 0.75, 0.75, 0.8, 0.75, 0.75, 0.75, 0.8], font_size=7.5)

    add_heading(doc, "四、零偏与噪声标准差")
    static = all_rows["static"]
    stat_table = [["项目", "均值", "零偏", "标准差", "说明"]]
    for column, note in [
        ("ax_g", "静止时 X 轴加速度"),
        ("ay_g", "静止时 Y 轴加速度"),
        ("az_g", "静止时 Z 轴加速度，理想约 1 g"),
        ("gx_dps", "陀螺 X 轴零偏"),
        ("gy_dps", "陀螺 Y 轴零偏"),
        ("gz_dps", "陀螺 Z 轴零偏"),
        ("pitch_dmp", "DMP Pitch 稳定性"),
        ("pitch_comp", "互补滤波 Pitch 稳定性"),
    ]:
        values = [float(row[column]) for row in static]
        avg = mean(values)
        zero = avg
        if column == "az_g":
            zero = avg - (1.0 if avg >= 0 else -1.0)
        stat_table.append([column, f"{avg:.6f}", f"{zero:.6f}", f"{stdev(values):.6f}", note])
    add_table(doc, stat_table, [0.95, 1.05, 1.05, 1.05, 2.65], font_size=8)

    add_heading(doc, "五、姿态曲线与方向验证")
    for key in ["static", "tilt", "shake"]:
        doc.add_picture(str(plot_paths[key]), width=Inches(6.3))
        add_para(doc, f"图：{DATASETS[key]['label']}状态下 DMP 与互补滤波 Pitch 曲线。", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    tilt_stats = stats(all_rows["tilt"], "pitch_dmp")
    add_para(doc, f"倾斜数据中，校正后的 pitch_dmp 范围为 {tilt_stats['min']:.2f}° ~ {tilt_stats['max']:.2f}°，覆盖前倾负角与后倾正角。原始采集方向与要求相反，因此采用统一输出符号校正，固件中对应设置 MPU_PITCH_OUTPUT_SIGN=-1。")

    doc.add_page_break()
    add_heading(doc, "六、DMP 与互补滤波对比")
    comp_table = [["维度", "DMP", "互补滤波"]]
    comp_table += [
        ["噪声大小", f"静止段标准差 {stats(static, 'pitch_dmp')['std']:.4f}°，曲线最平稳。", f"静止段标准差 {stats(static, 'pitch_comp')['std']:.4f}°，但存在固定角度偏置。"],
        ["响应延迟", "倾斜段变化较平滑，抗噪声强。", "依赖陀螺积分和加速度修正，动态响应快但可能出现超调。"],
        ["长期漂移", "DMP 内部融合和校准后长期稳定性较好。", "陀螺积分项会累积漂移，需靠加速度低频修正拉回。"],
    ]
    add_table(doc, comp_table, [1.0, 2.75, 2.75], font_size=8.5)

    add_heading(doc, "七、数据融合核心思想")
    add_para(doc, "加速度计可以通过重力方向得到低频姿态参考，长期不漂移，但在快速运动时容易受线性加速度干扰；陀螺仪可以测量角速度，短时间响应快、动态连续，但积分后会产生漂移。姿态融合的核心是利用陀螺仪承担短时动态响应，利用加速度计提供长期参考校正，从而同时降低噪声、延迟和漂移。")

    DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_OUT)


def register_pdf_fonts() -> str:
    font_path = Path("C:/Windows/Fonts/simsun.ttc")
    pdfmetrics.registerFont(TTFont("SimSun", str(font_path)))
    return "SimSun"


def pdf_table(data: list[list[str]], widths: list[float], font: str, font_size: int = 8) -> Table:
    table = Table(data, colWidths=[w * inch for w in widths], repeatRows=1)
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font),
        ("FONTSIZE", (0, 0), (-1, -1), font_size),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.black),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.black),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("FONTNAME", (0, 0), (-1, 0), font),
    ]))
    return table


def build_pdf(all_rows: dict[str, list[dict[str, float | str]]], plot_paths: dict[str, Path]) -> None:
    font = register_pdf_fonts()
    styles = getSampleStyleSheet()
    normal = ParagraphStyle("cn", parent=styles["Normal"], fontName=font, fontSize=9.5, leading=14, textColor=colors.black, spaceAfter=5)
    title = ParagraphStyle("title", parent=normal, fontSize=17, leading=22, alignment=TA_CENTER, spaceAfter=8)
    h1 = ParagraphStyle("h1", parent=normal, fontSize=13, leading=18, spaceBefore=9, spaceAfter=5)
    caption = ParagraphStyle("cap", parent=normal, fontSize=8.5, leading=12, alignment=TA_CENTER)

    story = [Paragraph("MPU6050 传感器通信与姿态解算实验报告", title), Paragraph("文件名占位：学号_姓名.pdf", caption)]
    story += [Paragraph("一、实验连接与采集设置", h1)]
    story.append(PdfImage(str(DIAGRAM_OUT), width=6.2 * inch, height=2.9 * inch))
    story.append(Paragraph("图：MPU6050 与 STM32 的 I2C、供电和串口连接。", caption))
    story.append(pdf_table([
        ["STM32 引脚", "MPU6050 引脚", "说明"],
        ["PB14", "SCL", "软件 I2C 时钟线"],
        ["PB15", "SDA", "软件 I2C 数据线"],
        ["3V3", "VCC", "传感器供电"],
        ["GND", "GND", "公共地"],
        ["PA9", "USART1_TX", "串口输出 115200 8N1"],
        ["PB9", "INT", "可选中断输入"],
    ], [1.2, 1.4, 3.4], font))
    story.append(Spacer(1, 6))
    story.append(Paragraph("采样周期约 20 ms，对应更新频率约 50 Hz。根据实测安装方向，车体前后倾对应 DMP 的 Roll 轴，报告中的 Pitch 已按统一符号校正处理。", normal))

    story.append(Paragraph("二、三种状态采集概况", h1))
    overview = [["状态", "文件", "样本", "时长(s)", "周期(ms)", "频率(Hz)", "Pitch DMP 范围(deg)"]]
    for key, meta in DATASETS.items():
        rows = all_rows[key]
        period = median_period_ms(rows)
        duration = (float(rows[-1]["ms"]) - float(rows[0]["ms"])) / 1000.0
        pitch = stats(rows, "pitch_dmp")
        overview.append([meta["label"], meta["file"].name, str(len(rows)), f"{duration:.2f}", f"{period:.1f}", f"{1000.0 / period:.1f}", f"{pitch['min']:.2f} ~ {pitch['max']:.2f}"])
    story.append(pdf_table(overview, [0.75, 0.85, 0.55, 0.65, 0.65, 0.65, 1.5], font, 7.5))

    story.append(PageBreak())
    story.append(Paragraph("三、静止原始数据记录", h1))
    raw_table = [["序号", "ms", "ax_raw", "ay_raw", "az_raw", "gx_raw", "gy_raw", "gz_raw", "pitch_dmp"]]
    for i, row in enumerate(all_rows["static"][:10], 1):
        raw_table.append([str(i), f"{float(row['ms']):.0f}", f"{float(row['ax_raw']):.0f}", f"{float(row['ay_raw']):.0f}", f"{float(row['az_raw']):.0f}", f"{float(row['gx_raw']):.0f}", f"{float(row['gy_raw']):.0f}", f"{float(row['gz_raw']):.0f}", f"{float(row['pitch_dmp']):.3f}"])
    story.append(pdf_table(raw_table, [0.38, 0.62, 0.65, 0.65, 0.7, 0.65, 0.65, 0.65, 0.8], font, 7))

    story.append(PageBreak())
    story.append(Paragraph("四、零偏与噪声标准差", h1))
    static = all_rows["static"]
    stat_table = [["项目", "均值", "零偏", "标准差", "说明"]]
    for column, note in [
        ("ax_g", "X 轴加速度"),
        ("ay_g", "Y 轴加速度"),
        ("az_g", "Z 轴加速度，理想约 1 g"),
        ("gx_dps", "陀螺 X 轴零偏"),
        ("gy_dps", "陀螺 Y 轴零偏"),
        ("gz_dps", "陀螺 Z 轴零偏"),
        ("pitch_dmp", "DMP Pitch 稳定性"),
        ("pitch_comp", "互补滤波 Pitch 稳定性"),
    ]:
        values = [float(row[column]) for row in static]
        avg = mean(values)
        zero = avg - 1.0 if column == "az_g" and avg >= 0 else avg
        stat_table.append([column, f"{avg:.6f}", f"{zero:.6f}", f"{stdev(values):.6f}", note])
    story.append(pdf_table(stat_table, [0.8, 0.9, 0.9, 0.9, 2.5], font, 7.5))

    story.append(Paragraph("五、姿态曲线与方向验证", h1))
    for key in ["static", "tilt", "shake"]:
        story.append(PdfImage(str(plot_paths[key]), width=6.2 * inch, height=3.5 * inch))
        story.append(Paragraph(f"图：{DATASETS[key]['label']}状态下 DMP 与互补滤波 Pitch 曲线。", caption))
    tilt_stats = stats(all_rows["tilt"], "pitch_dmp")
    story.append(Paragraph(f"倾斜数据中，校正后的 pitch_dmp 范围为 {tilt_stats['min']:.2f}° ~ {tilt_stats['max']:.2f}°，覆盖前倾负角与后倾正角。原始采集方向与要求相反，因此采用统一输出符号校正，固件中对应设置 MPU_PITCH_OUTPUT_SIGN=-1。", normal))

    story.append(PageBreak())
    story.append(Paragraph("六、DMP 与互补滤波对比", h1))
    comp_table = [
        ["维度", "DMP", "互补滤波"],
        ["噪声大小", f"静止段标准差 {stats(static, 'pitch_dmp')['std']:.4f}°，曲线最平稳。", f"静止段标准差 {stats(static, 'pitch_comp')['std']:.4f}°，但存在固定角度偏置。"],
        ["响应延迟", "倾斜段变化较平滑，抗噪声强。", "动态响应快，但可能出现超调。"],
        ["长期漂移", "内部融合和校准后长期稳定性较好。", "陀螺积分项会漂移，需加速度低频修正。"],
    ]
    story.append(pdf_table(comp_table, [0.8, 2.55, 2.55], font, 7.5))
    story.append(Paragraph("七、数据融合核心思想", h1))
    story.append(Paragraph("加速度计通过重力方向提供长期姿态参考，不易长期漂移，但快速运动时会受线性加速度干扰；陀螺仪测量角速度，短时响应快、动态连续，但积分后会产生漂移。融合的核心是用陀螺仪承担短时动态响应，用加速度计提供长期低频校正，从而兼顾噪声、延迟和漂移。", normal))

    PDF_OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(PDF_OUT), pagesize=letter, rightMargin=0.75 * inch, leftMargin=0.75 * inch, topMargin=0.75 * inch, bottomMargin=0.75 * inch)
    doc.build(story)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    PLOT_DIR.mkdir(parents=True, exist_ok=True)
    draw_connection_diagram(DIAGRAM_OUT)
    all_rows: dict[str, list[dict[str, float | str]]] = {}
    plot_paths: dict[str, Path] = {}
    for key, meta in DATASETS.items():
        rows = read_rows(meta["file"], key)
        if not rows:
            raise SystemExit(f"No valid rows found in {meta['file']}")
        all_rows[key] = rows
        plot_path = PLOT_DIR / f"{key}_pitch_compare.png"
        draw_plot(rows, PLOT_TITLES[key], plot_path)
        plot_paths[key] = plot_path
    build_docx(all_rows, plot_paths)
    build_pdf(all_rows, plot_paths)
    print(f"Wrote {DOCX_OUT}")
    print(f"Wrote {PDF_OUT}")


if __name__ == "__main__":
    main()
