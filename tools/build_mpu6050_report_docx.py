#!/usr/bin/env python3
"""Build a black-and-white DOCX MPU6050 experiment report template."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT = ROOT / "output" / "docx" / "学号_姓名.docx"
DEFAULT_ANALYSIS_DIR = ROOT / "output" / "mpu6050"
FONT_NAME = "SimSun"
HEADING_FONT = "SimHei"
TABLE_WIDTH_DXA = 9360


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    parser.add_argument("--analysis-dir", type=Path, default=DEFAULT_ANALYSIS_DIR)
    parser.add_argument("--student-id", default="学号")
    parser.add_argument("--student-name", default="姓名")
    return parser.parse_args()


def set_run_font(run, size: float | None = None, bold: bool | None = None, font=FONT_NAME) -> None:
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.color.rgb = None
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_paragraph_font(paragraph, size: float | None = None, bold: bool | None = None, font=FONT_NAME) -> None:
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, font=font)


def set_cell_text(cell, text: str, size: float = 9, bold: bool = False, align=WD_ALIGN_PARAGRAPH.CENTER) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def set_table_width(table, width_dxa=TABLE_WIDTH_DXA) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")


def set_column_widths(table, widths_in: list[float]) -> None:
    for row in table.rows:
        for index, width in enumerate(widths_in):
            cell = row.cells[index]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def format_table(table, widths_in: list[float] | None = None) -> None:
    set_table_width(table)
    set_table_borders(table)
    if widths_in:
        set_column_widths(table, widths_in)
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading("", level=level)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_run_font(run, size=16 if level == 1 else 12, bold=True, font=HEADING_FONT)


def add_body(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5)


def read_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(10.5)

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = HEADING_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), HEADING_FONT)
        style.font.color.rgb = None

    return doc


def add_title_page(doc: Document, student_id: str, student_name: str) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(140)
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run("MPU6050 传感器通信与姿态解算实验报告")
    set_run_font(run, size=22, bold=True, font=HEADING_FONT)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(48)
    run = subtitle.add_run("黑白简洁版")
    set_run_font(run, size=12)

    for label, value in (
        ("学号", student_id),
        ("姓名", student_name),
        ("平台", "STM32F103 + MPU6050"),
        ("通信", "软件 I2C，PB14=SCL，PB15=SDA"),
    ):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(1.5)
        paragraph.paragraph_format.space_after = Pt(10)
        run = paragraph.add_run(f"{label}：{value}")
        set_run_font(run, size=12)

    add_body(
        doc,
        "说明：本报告模板不包含伪造实验数据。请先烧录 MPU 实验模式固件，通过 USART1 采集 CSV，"
        "再运行 tools/mpu6050_analyze.py 生成统计表与曲线截图，最后重新生成报告。",
    )
    doc.add_page_break()


def add_connection_section(doc: Document) -> None:
    add_heading(doc, "1. MPU6050 与 STM32 的 I2C 连接")
    add_body(
        doc,
        "依据 C10B 主板原理图与工程引脚定义，MPU6050 使用 PB14/PB15 作为软件 I2C 引脚，"
        "电源为 3V3，地为 GND，PB9 可作为 INT 中断输入。",
    )
    add_body(doc, "串口输出：USART1，PA9=TX，PA10=RX，115200 8N1。")

    table = doc.add_table(rows=1, cols=3)
    format_table(table, [2.0, 2.5, 2.0])
    for cell, text in zip(table.rows[0].cells, ["STM32F103", "连接", "MPU6050"]):
        set_cell_text(cell, text, bold=True)
    for left, middle, right in [
        ("PB14", "->", "SCL"),
        ("PB15", "->", "SDA"),
        ("3V3", "->", "VCC"),
        ("GND", "->", "GND"),
        ("PB9", "->", "INT"),
    ]:
        cells = table.add_row().cells
        set_cell_text(cells[0], left)
        set_cell_text(cells[1], middle)
        set_cell_text(cells[2], right)
    doc.add_page_break()


def add_raw_data_section(doc: Document, analysis_dir: Path) -> None:
    add_heading(doc, "2. 原始数据读取与单位换算")
    add_body(
        doc,
        "采集要求：静止、倾斜、快速晃动三种状态均需记录。静止状态至少取 10 组有效样本，"
        "并计算零偏与噪声标准差。",
    )

    samples = read_csv(analysis_dir / "static_samples.csv")[:10]
    columns = ["ms", "ax_raw", "ay_raw", "az_raw", "gx_raw", "gy_raw", "gz_raw", "pitch_dmp"]
    headers = ["ms", "ax", "ay", "az", "gx", "gy", "gz", "Pitch"]
    table = doc.add_table(rows=1, cols=len(headers))
    format_table(table, [0.65, 0.8, 0.8, 0.8, 0.8, 0.8, 0.8, 1.05])
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_text(cell, header, bold=True)
    for index in range(10):
        row = table.add_row().cells
        if index < len(samples):
            data = [samples[index].get(column, "") for column in columns]
        else:
            data = ["待采集"] + [""] * (len(headers) - 1)
        for cell, text in zip(row, data):
            set_cell_text(cell, text, size=8)

    doc.add_paragraph()
    summary = read_csv(analysis_dir / "static_summary.csv")
    by_column = {row.get("column"): row for row in summary}
    stats = doc.add_table(rows=1, cols=4)
    format_table(stats, [1.6, 1.6, 1.6, 1.6])
    for cell, header in zip(stats.rows[0].cells, ["项目", "均值", "零偏", "标准差"]):
        set_cell_text(cell, header, bold=True)
    for column in ["ax_g", "ay_g", "az_g", "gx_dps", "gy_dps", "gz_dps", "pitch_dmp", "pitch_comp"]:
        row = stats.add_row().cells
        values = by_column.get(column)
        data = [column, "待计算", "待计算", "待计算"]
        if values:
            data = [column, values.get("mean", ""), values.get("zero_bias", ""), values.get("stddev", "")]
        for cell, text in zip(row, data):
            set_cell_text(cell, text, size=8.5)
    doc.add_page_break()


def add_comparison_section(doc: Document, analysis_dir: Path) -> None:
    add_heading(doc, "3. DMP 与互补滤波姿态解算对比")
    curve = analysis_dir / "pitch_compare.png"
    if curve.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(curve), width=Inches(6.2))
    else:
        table = doc.add_table(rows=1, cols=1)
        format_table(table, [6.5])
        set_cell_text(table.rows[0].cells[0], "待插入上位机姿态曲线截图", size=12)
        table.rows[0].height = Inches(2.0)

    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=3)
    format_table(table, [1.4, 2.55, 2.55])
    for cell, header in zip(table.rows[0].cells, ["维度", "DMP", "互补滤波"]):
        set_cell_text(cell, header, bold=True)
    rows = [
        ("噪声大小", "静止段更平滑，标准差较小。", "会受加速度瞬时噪声影响。"),
        ("响应延迟", "输出稳定，快速变化时略滞后。", "陀螺积分项响应较快。"),
        ("长期漂移", "内部融合与校准可抑制漂移。", "受陀螺零偏影响，长时可能漂移。"),
    ]
    for data in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, data):
            set_cell_text(cell, text, size=9, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_body(
        doc,
        "正负方向验证：按实验要求，前倾 Pitch 应为负，后倾 Pitch 应为正。若实物安装方向相反，"
        "在固件中通过 MPU_PITCH_OUTPUT_SIGN 统一修正输出方向。",
    )
    doc.add_page_break()


def add_fusion_section(doc: Document) -> None:
    add_heading(doc, "4. 数据融合核心思想")
    add_body(doc, "加速度计能通过重力方向估计倾角，低频稳定、长期不漂移，但在快速晃动或受到线加速度干扰时噪声明显。")
    add_body(doc, "陀螺仪测量角速度，积分后可得到角度，短时响应快、动态性能好，但零偏会随时间积分成漂移。")
    add_body(
        doc,
        "互补滤波把两者按频率特性融合：低频信任加速度计，高频信任陀螺仪。DMP 则在芯片内部完成更复杂的六轴融合与校准，"
        "主控只需读取四元数并换算欧拉角。",
    )
    add_heading(doc, "提交前检查", level=2)
    for item in [
        "[ ] 已采集 static、tilt、shake 三段 CSV 数据。",
        "[ ] 静止状态有效样本不少于 10 组。",
        "[ ] 已运行 tools/mpu6050_analyze.py 并生成统计表和曲线图。",
        "[ ] 已把占位学号、姓名替换为真实信息。",
        "[ ] DOCX 全文和表格为黑白，无伪造数据。",
    ]:
        add_body(doc, item)


def add_footer(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("MPU6050 姿态解算实验报告")
        set_run_font(run, size=8)


def main() -> int:
    args = parse_args()
    args.out.parent.mkdir(parents=True, exist_ok=True)
    doc = setup_document()
    add_title_page(doc, args.student_id, args.student_name)
    add_connection_section(doc)
    add_raw_data_section(doc, args.analysis_dir)
    add_comparison_section(doc, args.analysis_dir)
    add_fusion_section(doc)
    add_footer(doc)
    doc.save(args.out)
    print(args.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
